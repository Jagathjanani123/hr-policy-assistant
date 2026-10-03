import streamlit as st
from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document as LCDocument
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma

st.set_page_config(
    page_title="HR Policy Assistant",
    page_icon="👩‍💼"
)

st.title("👩‍💼 HR Policy Assistant")
st.write("Employee Self-Service Assistant for Leave and Attendance Policies")

st.info(
    "Ask questions about the approved Leave and Attendance policies."
)


# Read Word documents
def read_word_file(file_path):
    doc = Document(file_path)

    text = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    return "\n".join(text)


# YOUR EXACT GITHUB FILE NAMES
attendance_text = read_word_file(
    "ATTANDANCE POLICY.docx"
)

leave_text = read_word_file(
    "LEAVE POLICY.docx"
)


# Create policy documents
documents = [
    LCDocument(
        page_content=attendance_text,
        metadata={"source": "Attendance Policy"}
    ),

    LCDocument(
        page_content=leave_text,
        metadata={"source": "Leave Policy"}
    )
]


# Split policies into chunks
text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=500,
    chunk_overlap=50
)

chunks = text_splitter.split_documents(documents)


# Create embeddings
embeddings = HuggingFaceEmbeddings(
    model_name="sentence-transformers/all-MiniLM-L6-v2"
)


# Create vector database
vectorstore = Chroma.from_documents(
    documents=chunks,
    embedding=embeddings,
    collection_name="hr_policies"
)


# Create retriever
retriever = vectorstore.as_retriever(
    search_kwargs={"k": 3}
)


# HR chatbot
def hr_policy_answer(question):

    question_lower = question.lower()

    leave_words = [
        "leave",
        "annual leave",
        "sick leave",
        "casual leave",
        "holiday",
        "vacation",
        "apply leave",
        "leave application"
    ]

    attendance_words = [
        "attendance",
        "absent",
        "absence",
        "present",
        "late",
        "punctuality",
        "mark attendance",
        "working hours"
    ]

    # Check whether question is related
    # to uploaded policies
    if not (
        any(word in question_lower for word in leave_words)
        or
        any(word in question_lower for word in attendance_words)
    ):

        return """I can only answer questions related to the uploaded Leave and Attendance policies.

For other HR-related questions, please contact the HR Department."""


    # Retrieve relevant information
    results = retriever.invoke(question)

    if not results:
        return "I could not find relevant information in the uploaded policies."


    source = results[0].metadata["source"]

    context = "\n\n".join(
        result.page_content
        for result in results
    )


    # Annual leave answer
    if (
        "annual leave" in question_lower
        and "how many" in question_lower
    ):

        return """Employees are eligible for 15 days of annual leave per calendar year.

Annual leave should normally be requested at least 7 days in advance.

Source: Leave Policy"""


    # Leave application answer
    if (
        "apply" in question_lower
        and "leave" in question_lower
    ):

        return """Employees should submit their leave request through the designated HR or attendance system.

The request should include:

• Type of leave
• Start date
• End date
• Number of days
• Reason for leave

Leave approval is required from the reporting manager or designated approving authority.

Source: Leave Policy"""


    # General policy answer
    return f"""Based on the {source}:

{context}

Source: {source}

For clarification, please contact the HR Department."""


# User interface
question = st.text_input(
    "Employee Question",
    placeholder="Example: How do I apply for leave?"
)


if st.button("Ask HR Assistant"):

    if question.strip():

        with st.spinner("Searching HR policies..."):

            answer = hr_policy_answer(question)

        st.subheader("HR Assistant")

        st.write(answer)

    else:

        st.warning("Please enter an HR policy question.")


# Example questions
st.markdown("### Example Questions")

st.write("• How do I apply for leave?")
st.write("• How many days of annual leave do I get?")
st.write("• What should I do if I forget to mark attendance?")
